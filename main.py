import os
import re
import pandas as pd
import customtkinter as ctk
from tkinter import ttk, filedialog, messagebox
import unicodedata
from PIL import Image
from fpdf import FPDF
import datetime
import time
import xlrd
import openpyxl

def cargar_catalogo(ruta_archivo):
    try:
        return pd.read_excel(ruta_archivo)
    except Exception:
        if str(ruta_archivo).lower().endswith(('.xls', '.XLS')):
            try:
                import xlrd
                libro_rescatado = xlrd.open_workbook(ruta_archivo, ignore_workbook_corruption=True)
                return pd.read_excel(libro_rescatado)
            except Exception:
                return None
        return None

ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")

# ==========================================
# VENTANA EMERGENTE PARA AGREGAR / EDITAR
# ==========================================
class VentanaProducto(ctk.CTkToplevel):
    def __init__(self, parent, accion="nuevo", indice_editar=None, datos_actuales=None):
        super().__init__(parent)
        self.parent = parent
        self.accion = accion
        self.indice_editar = indice_editar
        
        titulo = "Agregar Nuevo Producto" if accion == "nuevo" else "Editar Producto"
        self.title(titulo)
        self.geometry("450x650")
        self.grab_set()
        
        ctk.CTkLabel(self, text=titulo, font=("Arial", 18, "bold")).pack(pady=10)

        self.lbl_img_editar = ctk.CTkLabel(self, text="")
        self.lbl_img_editar.pack(pady=5)

        if accion == "editar" and datos_actuales is not None:
            codigo_prod = str(datos_actuales.get("CODIGO", ""))
            ruta_jpg = f"fotos/{codigo_prod}.jpg"
            ruta_png = f"fotos/{codigo_prod}.png"
            img_cargada = None
            
            if os.path.exists(ruta_jpg): img_cargada = Image.open(ruta_jpg)
            elif os.path.exists(ruta_png): img_cargada = Image.open(ruta_png)

            if img_cargada:
                img_ctk = ctk.CTkImage(light_image=img_cargada, size=(120, 120))
                self.lbl_img_editar.configure(image=img_ctk, text="")
                self.lbl_img_editar.image = img_ctk
            else:
                self.lbl_img_editar.configure(text="📸 Sin foto asignada", text_color="gray", font=("Arial", 12, "italic"))
        else:
            self.lbl_img_editar.configure(text="📸 Sube una foto luego de guardar", text_color="gray", font=("Arial", 12, "italic"))
        
        self.entradas = {}
        self.campos_ingreso = ["CODIGO", "DESCRIPCION", "PRECIO COSTO", "PRECIO VENTA", "PROVEEDOR", "FECHA"]
        nombres_mostrar = ["Código", "Descripción", "Precio Costo (Sin IVA)", "Precio Venta", "Proveedor", "Fecha (ej. 2026-06-28)"]
        
        for i, campo in enumerate(self.campos_ingreso):
            ctk.CTkLabel(self, text=nombres_mostrar[i]).pack(anchor="w", padx=40)
            entrada = ctk.CTkEntry(self, width=370)
            entrada.pack(pady=5, padx=40)
            self.entradas[campo] = entrada
            
            if accion == "editar" and datos_actuales is not None:
                valor = str(datos_actuales.get(campo, ""))
                if valor != "nan" and valor != "":
                    entrada.insert(0, valor)
                    
            elif accion == "nuevo" and campo == "FECHA":
                fecha_hoy = datetime.datetime.now().strftime("%Y-%m-%d")
                entrada.insert(0, fecha_hoy)
                    
        ctk.CTkButton(self, text="💾 Guardar Cambios", command=self.guardar_datos).pack(pady=15)
        
    def guardar_datos(self):
        nuevos_datos = {}
        for campo in self.campos_ingreso:
            nuevos_datos[campo] = self.entradas[campo].get().strip()
            
        try:
            costo = float(nuevos_datos["PRECIO COSTO"]) if nuevos_datos["PRECIO COSTO"] else 0.0
            venta = float(nuevos_datos["PRECIO VENTA"]) if nuevos_datos["PRECIO VENTA"] else 0.0
        except ValueError:
            messagebox.showerror("Error", "Los precios deben ser números válidos.")
            return

        costo_iva = costo * 1.19
        ganancia = venta - costo_iva

        nuevos_datos["PRECIO COSTO"] = costo
        nuevos_datos["PRECIO C+IVA"] = costo_iva
        nuevos_datos["PRECIO VENTA"] = venta
        nuevos_datos["GANANCIA"] = ganancia

        if self.accion == "nuevo":
            codigo_ingresado = str(nuevos_datos["CODIGO"]).strip()
            df_actual = self.parent.catalogo_memoria
            
            if df_actual is not None and not df_actual.empty:
                coincidencias = df_actual[df_actual['CODIGO'].astype(str).str.strip() == codigo_ingresado]
                
                if not coincidencias.empty:
                    respuesta = messagebox.askyesnocancel(
                        "Código Repetido Detectado",
                        f"El código '{codigo_ingresado}' ya existe en el catálogo.\n\n[SÍ] para Sobrescribir el antiguo.\n[NO] para agregarlo como nuevo separado.\n[Cancelar] para no hacer nada."
                    )
                    
                    if respuesta is None:
                        return
                    elif respuesta is True:
                        indice_editar = coincidencias.index[0]
                        self.parent.actualizar_producto_memoria(indice_editar, nuevos_datos)
                        self.destroy()
                        return
            self.parent.agregar_producto_memoria(nuevos_datos)
        else:
            self.parent.actualizar_producto_memoria(self.indice_editar, nuevos_datos)
            
        self.destroy()

# ==========================================
# VENTANA EMERGENTE PARA PRODUCTO EXTRA
# ==========================================
class VentanaProductoExtra(ctk.CTkToplevel):
    def __init__(self, parent, destino="cotizacion"):
        super().__init__(parent)
        self.parent = parent
        self.destino = destino
        
        self.title("Agregar Ítem Fuera de Catálogo")
        self.geometry("400x350")
        self.grab_set()
        
        texto_titulo = "Ítem Especial (Solo para esta Cotización)" if destino == "cotizacion" else "Ítem Especial (Venta Rápida)"
        ctk.CTkLabel(self, text=texto_titulo, font=("Arial", 14, "bold")).pack(pady=15)
        
        ctk.CTkLabel(self, text="Descripción del Producto / Servicio:").pack(anchor="w", padx=30)
        self.ent_desc = ctk.CTkEntry(self, width=340)
        self.ent_desc.pack(pady=5, padx=30)
        
        ctk.CTkLabel(self, text="Precio Unitario (Venta con IVA):").pack(anchor="w", padx=30)
        self.ent_precio = ctk.CTkEntry(self, width=340)
        self.ent_precio.pack(pady=5, padx=30)

        ctk.CTkLabel(self, text="Cantidad:").pack(anchor="w", padx=30)
        self.ent_cant = ctk.CTkEntry(self, width=340)
        self.ent_cant.insert(0, "1")
        self.ent_cant.pack(pady=5, padx=30)
                    
        ctk.CTkButton(self, text="➕ Añadir", command=self.guardar_extra).pack(pady=20)
        
    def guardar_extra(self):
        desc = self.ent_desc.get().strip()
        try:
            precio = float(self.ent_precio.get())
            cant = int(self.ent_cant.get())
        except ValueError:
            messagebox.showerror("Error", "El precio y la cantidad deben ser números válidos.")
            return

        if not desc:
            messagebox.showwarning("Aviso", "La descripción es obligatoria.")
            return

        if self.destino == "cotizacion":
            self.parent.agregar_extra_cotizacion(desc, precio, cant)
        elif self.destino == "pos":
            self.parent.agregar_extra_pos(desc, precio, cant)
            
        self.destroy()

# ==========================================
# VENTANA PRINCIPAL
# ==========================================
class AppVentas(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Tienda de Instrumentos - Sistema de Ventas")
        self.geometry("1150x650") 
        
        self.catalogo_memoria = None
        self.carrito = {} 
        self.carrito_cotizacion = {}

        self.columnas_reales = ["CODIGO", "DESCRIPCION", "PRECIO COSTO", "PRECIO C+IVA", "PRECIO VENTA", "GANANCIA", "PROVEEDOR", "FECHA"]
        
        self.tabview = ctk.CTkTabview(self)
        self.tabview.pack(fill="both", expand=True, padx=10, pady=10)
        
        self.tab_catalogo = self.tabview.add("📦 Gestión de Catálogo")
        self.tab_ventas = self.tabview.add("🛒 Punto de Venta")
        self.tab_cotizaciones = self.tabview.add("📝 Cotizaciones")

        self.configurar_pestana_catalogo()
        self.configurar_pestana_ventas()
        self.configurar_pestana_cotizaciones()

        os.makedirs("fotos", exist_ok=True)
        self.cargar_archivo_base()

    def configurar_pestana_catalogo(self):
        self.frame_busqueda = ctk.CTkFrame(self.tab_catalogo, fg_color="transparent")
        self.frame_busqueda.pack(pady=5, padx=20, fill="x")
        
        self.entrada_busqueda = ctk.CTkEntry(self.frame_busqueda, placeholder_text="Buscar por código o descripción...", width=300)
        self.entrada_busqueda.pack(side="left", padx=5)
        self.entrada_busqueda.bind("<KeyRelease>", self.buscar_producto)
        
        self.boton_ver_todos = ctk.CTkButton(self.frame_busqueda, text="Ver Todos", width=90, fg_color="#2b7a4b", hover_color="#1e5c37", command=self.limpiar_busqueda)
        self.boton_ver_todos.pack(side="left", padx=5)

        self.opciones_orden = ctk.CTkOptionMenu(
            self.frame_busqueda,
            values=[
                "Predeterminado", 
                "Nombre: A-Z", 
                "Nombre: Z-A", 
                "Precio: Menor a Mayor", 
                "Precio: Mayor a Menor",
                "Fecha: Reciente a Antiguo",
                "Fecha: Antiguo a Reciente"
            ],
            command=self.ordenar_catalogo,
            width=210,
            fg_color="#3e4a59",
            button_color="#2b3440",
            button_hover_color="#1e252e"
        )
        self.opciones_orden.pack(side="left", padx=10)

        self.boton_exportar = ctk.CTkButton(self.frame_busqueda, text="📤 Exportar Respaldo", width=140, fg_color="#3e4a59", hover_color="#2b3440", command=self.exportar_excel)
        self.boton_exportar.pack(side="right", padx=5)

        self.boton_cargar_excels = ctk.CTkButton(self.frame_busqueda, text="📂 Cargar Excel", width=120, fg_color="#b85c00", hover_color="#8f4700", command=self.cargar_excels_manual)
        self.boton_cargar_excels.pack(side="right", padx=5)

        self.frame_acciones = ctk.CTkFrame(self.tab_catalogo, fg_color="transparent")
        self.frame_acciones.pack(pady=5, padx=20, fill="x")
        
        ctk.CTkButton(self.frame_acciones, text="➕ Nuevo Producto", width=120, fg_color="#225b82", command=self.abrir_ventana_nuevo).pack(side="left", padx=5)
        ctk.CTkButton(self.frame_acciones, text="✏️ Editar Seleccionado", width=140, fg_color="#6e5c00", hover_color="#544600", command=self.abrir_ventana_editar).pack(side="left", padx=5)
        ctk.CTkButton(self.frame_acciones, text="🗑️ Eliminar Seleccionado", width=150, fg_color="#8c1f1f", hover_color="#6b1616", command=self.eliminar_seleccionado).pack(side="left", padx=5)
        ctk.CTkButton(self.frame_acciones, text="📸 Asignar Foto", width=130, fg_color="#4b2b7a", hover_color="#371e5c", command=self.asignar_foto_seleccionado).pack(side="left", padx=5)

        self.btn_vaciar_catalogo = ctk.CTkButton(
            self.frame_acciones, text="⚠️ Vaciar Todo", width=100, 
            fg_color="transparent", border_width=1, border_color="#8c1f1f", text_color="#ff4d4d", hover_color="#3d1414", 
            command=self.vaciar_catalogo
        )
        self.btn_vaciar_catalogo.pack(side="right", padx=10)
        style = ttk.Style()
        style.theme_use("clam")
        style.configure("Treeview", background="#2a2d2e", foreground="white", rowheight=30, fieldbackground="#2a2d2e", borderwidth=0)
        style.map('Treeview', 
                  background=[('selected', '#1f538d')],
                  foreground=[('!selected', 'white'), ('selected', 'white')])
        style.configure("Treeview.Heading", background="#1f538d", foreground="white", font=("Arial", 10, "bold"))
        
        self.frame_tabla = ctk.CTkFrame(self.tab_catalogo)
        self.frame_tabla.pack(pady=10, padx=20, fill="both", expand=True)
        
        self.tabla = ttk.Treeview(self.frame_tabla, columns=self.columnas_reales, show="headings")
        
        nombres_cols = ["Código", "Descripción", "P. Costo", "P. Costo+IVA", "P. Venta", "Ganancia", "Proveedor", "Fecha"]
        anchos = [70, 300, 70, 80, 70, 70, 90, 70]
        
        for col, nombre, ancho in zip(self.columnas_reales, nombres_cols, anchos):
            self.tabla.heading(col, text=nombre)
            alineacion = "e" if "PRECIO" in col or "GANANCIA" in col else "center"
            if col == "DESCRIPCION": alineacion = "w"
            self.tabla.column(col, width=ancho, anchor=alineacion)
            
        self.tabla.pack(fill="both", expand=True, padx=2, pady=2)

    def configurar_pestana_ventas(self):
        self.frame_izq = ctk.CTkFrame(self.tab_ventas, fg_color="transparent")
        self.frame_izq.pack(side="left", fill="both", expand=True, padx=10, pady=10)
        
        self.frame_der = ctk.CTkFrame(self.tab_ventas)
        self.frame_der.pack(side="right", fill="both", expand=True, padx=10, pady=10)

        ctk.CTkLabel(self.frame_izq, text="🔍 Buscar Producto para Vender", font=("Arial", 16, "bold")).pack(anchor="w", pady=(0, 10))
        
        self.entrada_busqueda_pos = ctk.CTkEntry(self.frame_izq, placeholder_text="Escribe código o descripción...", width=350)
        self.entrada_busqueda_pos.pack(anchor="w", pady=5)
        self.entrada_busqueda_pos.bind("<KeyRelease>", self.buscar_producto_pos)
        
        self.frame_tabla_pos = ctk.CTkFrame(self.frame_izq)
        self.frame_tabla_pos.pack(fill="both", expand=True, pady=10)
        
        cols_pos = ("CODIGO", "DESCRIPCION", "PRECIO VENTA")
        self.tabla_pos = ttk.Treeview(self.frame_tabla_pos, columns=cols_pos, show="headings", height=8)
        self.tabla_pos.heading("CODIGO", text="Código")
        self.tabla_pos.heading("DESCRIPCION", text="Descripción")
        self.tabla_pos.heading("PRECIO VENTA", text="P. Venta")
        self.tabla_pos.column("CODIGO", width=60, anchor="center")
        self.tabla_pos.column("DESCRIPCION", width=220, anchor="w")
        self.tabla_pos.column("PRECIO VENTA", width=80, anchor="e")
        self.tabla_pos.pack(fill="both", expand=True)

        self.tabla_pos.bind("<Double-1>", self.agregar_al_carrito)

        self.btn_agregar_carrito = ctk.CTkButton(self.frame_izq, text="⬇️ Agregar al Carrito", fg_color="#2b7a4b", hover_color="#1e5c37", command=self.agregar_al_carrito)
        self.btn_agregar_carrito.pack(pady=10)

        self.btn_extra_pos = ctk.CTkButton(self.frame_izq, text="➕ Agregar Ítem Extra", fg_color="#4b2b7a", hover_color="#371e5c", command=self.abrir_ventana_extra_pos)
        self.btn_extra_pos.pack(pady=5)

        self.frame_imagen = ctk.CTkFrame(self.frame_der, height=150)
        self.frame_imagen.pack(fill="x", padx=10, pady=10)
        self.label_imagen = ctk.CTkLabel(self.frame_imagen, text="[ Imagen del Producto Aparecerá Aquí ]", text_color="gray", font=("Arial", 12, "italic"))
        self.label_imagen.pack(pady=50)

        ctk.CTkLabel(self.frame_der, text="🛒 Lista de Compras", font=("Arial", 16, "bold")).pack(anchor="w", padx=10)
        
        self.frame_carrito = ctk.CTkFrame(self.frame_der)
        self.frame_carrito.pack(fill="both", expand=True, padx=10, pady=5)
        
        cols_car = ("producto", "cant", "subtotal")
        self.tabla_carrito = ttk.Treeview(self.frame_carrito, columns=cols_car, show="headings", height=5)
        self.tabla_carrito.heading("producto", text="Producto")
        self.tabla_carrito.heading("cant", text="Cant.")
        self.tabla_carrito.heading("subtotal", text="Subtotal")
        self.tabla_carrito.column("producto", width=180)
        self.tabla_carrito.column("cant", width=50, anchor="center")
        self.tabla_carrito.column("subtotal", width=80, anchor="e")
        self.tabla_carrito.pack(fill="both", expand=True)

        frame_cantidades = ctk.CTkFrame(self.frame_der, fg_color="transparent")
        frame_cantidades.pack(fill="x", padx=10, pady=5)
        ctk.CTkButton(frame_cantidades, text="➕", width=40, command=lambda: self.modificar_cantidad("sumar")).pack(side="left", padx=5)
        ctk.CTkButton(frame_cantidades, text="➖", width=40, command=lambda: self.modificar_cantidad("restar")).pack(side="left", padx=5)
        ctk.CTkButton(frame_cantidades, text="🗑️ Quitar", width=80, fg_color="#8c1f1f", command=lambda: self.modificar_cantidad("quitar")).pack(side="left", padx=10)

        self.lbl_subtotal_pos = ctk.CTkLabel(self.frame_der, text="Neto: $ 0", font=("Arial", 14))
        self.lbl_subtotal_pos.pack(anchor="e", padx=20, pady=(10, 0))
        self.lbl_iva_pos = ctk.CTkLabel(self.frame_der, text="IVA (19%): $ 0", font=("Arial", 14))
        self.lbl_iva_pos.pack(anchor="e", padx=20)
        self.label_total = ctk.CTkLabel(self.frame_der, text="TOTAL: $ 0", font=("Arial", 28, "bold"), text_color="#45b56f")
        self.label_total.pack(anchor="e", padx=20, pady=(5, 10))

        self.btn_limpiar_venta = ctk.CTkButton(self.frame_der, text="🧹 Nueva Venta / Limpiar", fg_color="#b85c00", hover_color="#8f4700", command=self.limpiar_carrito)
        self.btn_limpiar_venta.pack(pady=10, fill="x", padx=20)
    
    def configurar_pestana_cotizaciones(self):
        self.frame_cot_izq = ctk.CTkFrame(self.tab_cotizaciones, fg_color="transparent")
        self.frame_cot_izq.pack(side="left", fill="both", expand=True, padx=10, pady=10)
        
        self.frame_cot_der = ctk.CTkFrame(self.tab_cotizaciones)
        self.frame_cot_der.pack(side="right", fill="both", expand=True, padx=10, pady=10)

        frame_cliente = ctk.CTkFrame(self.frame_cot_izq)
        frame_cliente.pack(fill="x", pady=(0, 10))
        ctk.CTkLabel(frame_cliente, text="👤 Datos del Cliente", font=("Arial", 14, "bold")).pack(pady=5)

        self.entradas_cliente = {}
        campos = ["Nombre / Institución *", "RUT", "Dirección", "Teléfono", "Correo"]
        for campo in campos:
            row = ctk.CTkFrame(frame_cliente, fg_color="transparent")
            row.pack(fill="x", padx=10, pady=2)
            ctk.CTkLabel(row, text=campo+":", width=140, anchor="w").pack(side="left")
            ent = ctk.CTkEntry(row, width=220)
            ent.pack(side="left", fill="x", expand=True)
            self.entradas_cliente[campo] = ent

        ctk.CTkLabel(self.frame_cot_izq, text="🔍 Buscar Producto para Cotizar", font=("Arial", 14, "bold")).pack(anchor="w", pady=(10, 5))
        self.entrada_busqueda_cot = ctk.CTkEntry(self.frame_cot_izq, placeholder_text="Escribe código o descripción...")
        self.entrada_busqueda_cot.pack(fill="x", pady=5)
        self.entrada_busqueda_cot.bind("<KeyRelease>", self.buscar_producto_cot)
        
        cols_cot = ("CODIGO", "DESCRIPCION", "PRECIO VENTA")
        self.tabla_cot_pos = ttk.Treeview(self.frame_cot_izq, columns=cols_cot, show="headings", height=5)
        self.tabla_cot_pos.heading("CODIGO", text="Código")
        self.tabla_cot_pos.heading("DESCRIPCION", text="Descripción")
        self.tabla_cot_pos.heading("PRECIO VENTA", text="P. Venta")
        self.tabla_cot_pos.column("CODIGO", width=60, anchor="center")
        self.tabla_cot_pos.column("DESCRIPCION", width=200, anchor="w")
        self.tabla_cot_pos.column("PRECIO VENTA", width=80, anchor="e")
        self.tabla_cot_pos.pack(fill="both", expand=True)
        self.tabla_cot_pos.bind("<Double-1>", self.agregar_al_carrito_cot)

        frame_botones_cot = ctk.CTkFrame(self.frame_cot_izq, fg_color="transparent")
        frame_botones_cot.pack(pady=10)
        
        ctk.CTkButton(frame_botones_cot, text="⬇️ Agregar del Catálogo", fg_color="#2b7a4b", hover_color="#1e5c37", command=self.agregar_al_carrito_cot).pack(side="left", padx=5)
        ctk.CTkButton(frame_botones_cot, text="➕ Agregar Ítem Extra", fg_color="#4b2b7a", hover_color="#371e5c", command=self.abrir_ventana_extra_cot).pack(side="left", padx=5)

        ctk.CTkLabel(self.frame_cot_der, text="📝 Detalle de Cotización", font=("Arial", 16, "bold")).pack(pady=10)
        
        self.tabla_carrito_cot = ttk.Treeview(self.frame_cot_der, columns=("producto", "cant", "subtotal"), show="headings", height=8)
        self.tabla_carrito_cot.heading("producto", text="Producto")
        self.tabla_carrito_cot.heading("cant", text="Cant.")
        self.tabla_carrito_cot.heading("subtotal", text="Subtotal")
        self.tabla_carrito_cot.column("producto", width=220)
        self.tabla_carrito_cot.column("cant", width=50, anchor="center")
        self.tabla_carrito_cot.column("subtotal", width=80, anchor="e")
        self.tabla_carrito_cot.pack(fill="both", expand=True, padx=10)

        frame_cant_cot = ctk.CTkFrame(self.frame_cot_der, fg_color="transparent")
        frame_cant_cot.pack(fill="x", padx=10, pady=5)
        ctk.CTkButton(frame_cant_cot, text="➕", width=40, command=lambda: self.modificar_cantidad_cot("sumar")).pack(side="left", padx=5)
        ctk.CTkButton(frame_cant_cot, text="➖", width=40, command=lambda: self.modificar_cantidad_cot("restar")).pack(side="left", padx=5)
        ctk.CTkButton(frame_cant_cot, text="🗑️ Quitar", width=80, fg_color="#8c1f1f", command=lambda: self.modificar_cantidad_cot("quitar")).pack(side="left", padx=10)

        self.lbl_subtotal_cot = ctk.CTkLabel(self.frame_cot_der, text="Neto: $ 0", font=("Arial", 14))
        self.lbl_subtotal_cot.pack(anchor="e", padx=20, pady=(10, 0))
        self.lbl_iva_cot = ctk.CTkLabel(self.frame_cot_der, text="IVA (19%): $ 0", font=("Arial", 14))
        self.lbl_iva_cot.pack(anchor="e", padx=20)
        self.lbl_total_cot = ctk.CTkLabel(self.frame_cot_der, text="TOTAL: $ 0", font=("Arial", 24, "bold"), text_color="#45b56f")
        self.lbl_total_cot.pack(anchor="e", padx=20, pady=(5, 15))

        self.btn_generar_cot = ctk.CTkButton(self.frame_cot_der, text="📄 GENERAR PDF", height=40, font=("Arial", 14, "bold"), command=self.generar_pdf_cotizacion)
        self.btn_generar_cot.pack(fill="x", padx=20, pady=5)
        
        self.btn_limpiar_cot = ctk.CTkButton(self.frame_cot_der, text="🧹 Limpiar Cotización", fg_color="#b85c00", hover_color="#8f4700", command=self.limpiar_cotizacion)
        self.btn_limpiar_cot.pack(fill="x", padx=20, pady=5)

    def actualizar_tabla_pos(self, dataframe_a_mostrar):
        for item in self.tabla_pos.get_children():
            self.tabla_pos.delete(item)
            
        if dataframe_a_mostrar is not None:
            df_mostrar = dataframe_a_mostrar[["CODIGO", "DESCRIPCION", "PRECIO VENTA"]].fillna("")
            for i, (indice, fila) in enumerate(df_mostrar.iterrows()):
                valores = fila.tolist()
                try:
                    if valores[2] != "": valores[2] = f"${float(valores[2]):,.0f}"
                except ValueError:
                    pass
                self.tabla_pos.insert("", "end", iid=str(indice), values=valores)
                if i % 50 == 0:
                    self.update_idletasks()

    def buscar_producto_pos(self, event=None):
        if self.catalogo_memoria is None: return
        palabra = self.entrada_busqueda_pos.get().lower().strip()
        if palabra == "":
            self.actualizar_tabla_pos(self.catalogo_memoria)
            return
        palabra_limpia = ''.join(c for c in unicodedata.normalize('NFKD', palabra) if unicodedata.category(c) != 'Mn')
        filtro = self.catalogo_memoria['columna_busqueda'].astype(str).str.contains(palabra_limpia, na=False)
        self.actualizar_tabla_pos(self.catalogo_memoria[filtro])

    def agregar_al_carrito(self, event=None):
        seleccion = self.tabla_pos.selection()
        if not seleccion:
            if event is not None: return
            messagebox.showwarning("Aviso", "Selecciona un producto para vender.")
            return
            
        indice_producto = seleccion[0] 
        
        if indice_producto in self.carrito:
            self.carrito[indice_producto]["cantidad"] += 1
        else:
            datos_producto = self.catalogo_memoria.loc[int(indice_producto)]
            self.carrito[indice_producto] = {
                "nombre": str(datos_producto["DESCRIPCION"]),
                "precio": float(datos_producto["PRECIO VENTA"]) if pd.notna(datos_producto["PRECIO VENTA"]) else 0.0,
                "cantidad": 1
            }
            
        codigo_prod = str(self.catalogo_memoria.loc[int(indice_producto), "CODIGO"])
        self.actualizar_vista_carrito(indice_seleccionado=indice_producto)
        self.mostrar_imagen_temporal(codigo_prod, self.carrito[indice_producto]["nombre"])

    def modificar_cantidad(self, accion):
        seleccion = self.tabla_carrito.selection()
        if not seleccion:
            messagebox.showwarning("Aviso", "Selecciona un producto del carrito.")
            return
            
        indice = seleccion[0]
        if accion == "sumar":
            self.carrito[indice]["cantidad"] += 1
        elif accion == "restar":
            self.carrito[indice]["cantidad"] -= 1
            if self.carrito[indice]["cantidad"] <= 0:
                del self.carrito[indice]
                indice = None 
        elif accion == "quitar":
            del self.carrito[indice]
            indice = None 
            
        self.actualizar_vista_carrito(indice_seleccionado=indice)

    def actualizar_vista_carrito(self, indice_seleccionado=None):
        for item in self.tabla_carrito.get_children():
            self.tabla_carrito.delete(item)
            
        total_venta = 0.0
        for indice, item in self.carrito.items():
            subtotal = item["precio"] * item["cantidad"]
            total_venta += subtotal
            self.tabla_carrito.insert("", "end", iid=indice, values=(item["nombre"], item["cantidad"], f"${subtotal:,.0f}"))
            
        neto = total_venta / 1.19
        iva = total_venta - neto
        
        self.lbl_subtotal_pos.configure(text=f"Neto: $ {neto:,.0f}")
        self.lbl_iva_pos.configure(text=f"IVA (19%): $ {iva:,.0f}")
        self.label_total.configure(text=f"TOTAL: $ {total_venta:,.0f}")
        
        if indice_seleccionado and indice_seleccionado in self.carrito:
            self.tabla_carrito.selection_set(indice_seleccionado)

        if not self.carrito:
            self.label_imagen.configure(image="", text="[ Imagen del Producto Aparecerá Aquí ]", text_color="gray")

    def limpiar_carrito(self):
        if not self.carrito: return 
        if messagebox.askyesno("Confirmar", "¿Deseas limpiar todo el carrito?"):
            self.carrito.clear() 
            self.actualizar_vista_carrito() 
            self.label_imagen.configure(image="", text="[ Imagen del Producto Aparecerá Aquí ]", text_color="gray")

    def mostrar_imagen_temporal(self, codigo_producto, nombre_producto):
        ruta_jpg = f"fotos/{codigo_producto}.jpg"
        ruta_png = f"fotos/{codigo_producto}.png"
        imagen_cargada = None
        
        if os.path.exists(ruta_jpg): imagen_cargada = Image.open(ruta_jpg)
        elif os.path.exists(ruta_png): imagen_cargada = Image.open(ruta_png)
            
        if imagen_cargada:
            img_ctk = ctk.CTkImage(light_image=imagen_cargada, size=(130, 130))
            self.label_imagen.configure(image=img_ctk, text="") 
            self.label_imagen.image = img_ctk 
        else:
            self.label_imagen.configure(image="", text=f"📸 Sin foto para:\n{nombre_producto}", text_color="#6eade8")

    def guardar_base_interna(self):
        if self.catalogo_memoria is not None and not self.catalogo_memoria.empty:
            df_guardar = self.catalogo_memoria[self.columnas_reales]
            try:
                df_guardar.to_excel("catalogo.xlsx", index=False)
            except Exception as e:
                print(f"Error al guardar: {e}")

    def exportar_excel(self):
        if self.catalogo_memoria is None or self.catalogo_memoria.empty: return
        ruta = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx *.xls")], initialfile="Respaldo_LeonardMusic.xlsx")
        if ruta:
            try:
                self.catalogo_memoria[self.columnas_reales].to_excel(ruta, index=False)
                messagebox.showinfo("Éxito", "Respaldo exportado.")
            except Exception as e:
                messagebox.showerror("Error", f"Error: {e}")

    def generar_columna_busqueda(self):
        df_seguro = self.catalogo_memoria.fillna("")
        busqueda_cruda = (df_seguro['CODIGO'].astype(str) + " " + df_seguro['DESCRIPCION'].astype(str)).str.lower()
        self.catalogo_memoria['columna_busqueda'] = busqueda_cruda.str.normalize('NFKD').str.encode('ascii', errors='ignore').str.decode('utf-8')

    def cargar_archivo_base(self):
        df = cargar_catalogo("catalogo.xlsx")
        if df is not None: self.procesar_y_mostrar_dataframe(df)

    def cargar_excels_manual(self):
        rutas = filedialog.askopenfilenames(title="Selecciona Excels", filetypes=[("Excel", "*.xlsx *.xls")])
        if not rutas: return
        
        lista_dfs = [cargar_catalogo(r) for r in rutas]
        lista_dfs = [df for df in lista_dfs if df is not None]
        
        if lista_dfs:
            df_nuevo_importado = pd.concat(lista_dfs, ignore_index=True)
            
            if self.catalogo_memoria is not None and not self.catalogo_memoria.empty:
                codigos_viejos = self.catalogo_memoria['CODIGO'].astype(str).str.strip().tolist()
                codigos_nuevos = df_nuevo_importado['CODIGO'].astype(str).str.strip().tolist()
                
                repetidos = set(codigos_viejos).intersection(set(codigos_nuevos))
                repetidos = {c for c in repetidos if c.lower() != 'nan' and c != ''}
                
                if repetidos:
                    respuesta = messagebox.askyesnocancel(
                        "Actualización Masiva Detectada", 
                        f"El Excel contiene {len(repetidos)} códigos que ya existen en el catálogo actual.\n\n[SÍ] Reemplazar los productos antiguos por los nuevos del Excel.\n[NO] Agregar todo sin borrar nada (crear duplicados).\n[Cancelar] Abortar la carga."
                    )
                    
                    if respuesta is None:
                        return
                    elif respuesta is True:
                        self.catalogo_memoria = self.catalogo_memoria[~self.catalogo_memoria['CODIGO'].astype(str).str.strip().isin(repetidos)]
            
            if self.catalogo_memoria is not None and not self.catalogo_memoria.empty:
                df_unificado = pd.concat([self.catalogo_memoria, df_nuevo_importado], ignore_index=True)
            else:
                df_unificado = df_nuevo_importado
                
            self.procesar_y_mostrar_dataframe(df_unificado)
            self.guardar_base_interna()

    def procesar_y_mostrar_dataframe(self, df_crudo):
        df_limpio = df_crudo.dropna(subset=['CODIGO'])
        df_limpio = df_limpio[df_limpio['CODIGO'].astype(str).str.strip() != '']
        df_limpio = df_limpio[df_limpio['CODIGO'].astype(str).str.lower() != 'nan']
        self.catalogo_memoria = df_limpio.reset_index(drop=True)

        self.generar_columna_busqueda()
        self.limpiar_busqueda()
        self.actualizar_tabla_pos(self.catalogo_memoria)
        self.actualizar_tabla_cot(self.catalogo_memoria)

    def actualizar_tabla(self, df_mostrar):
        for item in self.tabla.get_children(): self.tabla.delete(item)
        if df_mostrar is not None:
            df_limpio = df_mostrar[self.columnas_reales].fillna("")
            for i, (indice, fila) in enumerate(df_limpio.iterrows()):
                valores = fila.tolist()
                for idx in [2, 3, 4, 5]:
                    try:
                        if valores[idx] != "": valores[idx] = f"${float(valores[idx]):,.0f}"
                    except ValueError:
                        pass
                
                if " 00:00:00" in str(valores[7]):
                    valores[7] = str(valores[7]).replace(" 00:00:00", "")

                self.tabla.insert("", "end", iid=str(indice), values=valores)
                if i % 50 == 0:
                    self.update_idletasks()

    def buscar_producto(self, event=None):
        if self.catalogo_memoria is None: return
        palabra = self.entrada_busqueda.get().lower().strip()
        if palabra == "":
            self.actualizar_tabla(self.catalogo_memoria)
            return
        palabra_limpia = ''.join(c for c in unicodedata.normalize('NFKD', palabra) if unicodedata.category(c) != 'Mn')
        filtro = self.catalogo_memoria['columna_busqueda'].astype(str).str.contains(palabra_limpia, na=False)
        self.actualizar_tabla(self.catalogo_memoria[filtro])

    def limpiar_busqueda(self):
        self.entrada_busqueda.delete(0, 'end')
        self.actualizar_tabla(self.catalogo_memoria)

    def abrir_ventana_nuevo(self):
        if self.catalogo_memoria is None:
            self.catalogo_memoria = pd.DataFrame(columns=self.columnas_reales)
        VentanaProducto(self, accion="nuevo")

    def abrir_ventana_editar(self):
        seleccion = self.tabla.selection()
        if not seleccion: return
        indice = int(seleccion[0])
        datos_actuales = self.catalogo_memoria.loc[indice]
        VentanaProducto(self, accion="editar", indice_editar=indice, datos_actuales=datos_actuales)

    def eliminar_seleccionado(self):
        seleccion = self.tabla.selection()
        if not seleccion: return
        if messagebox.askyesno("Confirmar", "¿Eliminar producto?"):
            indice = int(seleccion[0])
            self.catalogo_memoria = self.catalogo_memoria.drop(indice)
            self.generar_columna_busqueda()
            
            self.entrada_busqueda.delete(0, 'end')
            self.ordenar_catalogo(self.opciones_orden.get())
            
            self.actualizar_tabla_pos(self.catalogo_memoria)
            self.actualizar_tabla_cot(self.catalogo_memoria)
            self.guardar_base_interna()

    def vaciar_catalogo(self):
        if self.catalogo_memoria is None or self.catalogo_memoria.empty:
            return
            
        confirmacion = messagebox.askyesno(
            "¡ADVERTENCIA CRÍTICA!",
            "¿Estás completamente seguro de que deseas VACIAR TODO el catálogo?\n\nEsto eliminará todos los productos de la pantalla y del programa.\n\n¡Esta acción no se puede deshacer!"
        )
        
        if confirmacion:
            self.catalogo_memoria = pd.DataFrame(columns=self.columnas_reales)
            self.generar_columna_busqueda()
            
            self.entrada_busqueda.delete(0, 'end')
            self.actualizar_tabla(self.catalogo_memoria)
            self.actualizar_tabla_pos(self.catalogo_memoria)
            self.actualizar_tabla_cot(self.catalogo_memoria)
            
            self.guardar_base_interna()
            messagebox.showinfo("Catálogo Vaciado", "El catálogo ha sido borrado con éxito. El programa está limpio.")

    def agregar_producto_memoria(self, nuevos_datos):
        df_nuevo = pd.DataFrame([nuevos_datos])
        self.catalogo_memoria = pd.concat([self.catalogo_memoria, df_nuevo], ignore_index=True)
        self.generar_columna_busqueda()
        
        self.entrada_busqueda.delete(0, 'end')
        self.ordenar_catalogo(self.opciones_orden.get())
        
        self.actualizar_tabla_pos(self.catalogo_memoria)
        self.actualizar_tabla_cot(self.catalogo_memoria)
        self.guardar_base_interna()
        
    def actualizar_producto_memoria(self, indice, nuevos_datos):
        for columna, valor in nuevos_datos.items():
            self.catalogo_memoria.at[indice, columna] = valor
        self.generar_columna_busqueda()
        
        self.entrada_busqueda.delete(0, 'end')
        self.ordenar_catalogo(self.opciones_orden.get())
        
        self.actualizar_tabla_pos(self.catalogo_memoria)
        self.actualizar_tabla_cot(self.catalogo_memoria)
        self.guardar_base_interna()

    def asignar_foto_seleccionado(self):
        seleccion = self.tabla.selection()
        if not seleccion:
            messagebox.showwarning("Aviso", "Selecciona un producto del catálogo para asignarle una foto.")
            return
            
        indice = int(seleccion[0])
        codigo = str(self.catalogo_memoria.loc[indice, "CODIGO"])
        
        ruta_imagen = filedialog.askopenfilename(
            title=f"Seleccionar foto para {codigo}",
            filetypes=[("Imágenes", "*.jpg *.jpeg *.png")]
        )
        
        if ruta_imagen:
            try:
                img = Image.open(ruta_imagen)
                if img.mode != 'RGB':
                    img = img.convert('RGB')
                
                img.save(f"fotos/{codigo}.jpg")
                messagebox.showinfo("Éxito", f"¡Foto asignada y guardada correctamente para el producto {codigo}!")
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo procesar la foto: {e}")

    def ordenar_catalogo(self, seleccion):
        if self.catalogo_memoria is None or self.catalogo_memoria.empty:
            return

        if seleccion == "Nombre: A-Z":
            self.catalogo_memoria = self.catalogo_memoria.sort_values(
                by="DESCRIPCION", ascending=True, key=lambda col: col.astype(str).str.lower()
            )
        elif seleccion == "Nombre: Z-A":
            self.catalogo_memoria = self.catalogo_memoria.sort_values(
                by="DESCRIPCION", ascending=False, key=lambda col: col.astype(str).str.lower()
            )
        elif seleccion == "Precio: Menor a Mayor":
            self.catalogo_memoria["PRECIO VENTA"] = pd.to_numeric(self.catalogo_memoria["PRECIO VENTA"], errors='coerce').fillna(0)
            self.catalogo_memoria = self.catalogo_memoria.sort_values(by="PRECIO VENTA", ascending=True)
        elif seleccion == "Precio: Mayor a Menor":
            self.catalogo_memoria["PRECIO VENTA"] = pd.to_numeric(self.catalogo_memoria["PRECIO VENTA"], errors='coerce').fillna(0)
            self.catalogo_memoria = self.catalogo_memoria.sort_values(by="PRECIO VENTA", ascending=False)
        elif seleccion == "Fecha: Reciente a Antiguo":
            self.catalogo_memoria = self.catalogo_memoria.sort_values(
                by="FECHA", ascending=False, key=lambda col: col.astype(str).str.lower()
            )
        elif seleccion == "Fecha: Antiguo a Reciente":
            self.catalogo_memoria = self.catalogo_memoria.sort_values(
                by="FECHA", ascending=True, key=lambda col: col.astype(str).str.lower()
            )
        elif seleccion == "Predeterminado":
            self.catalogo_memoria = self.catalogo_memoria.sort_values(
                by="CODIGO", ascending=True, key=lambda col: col.astype(str).str.lower()
            )
        
        self.buscar_producto()

    def actualizar_tabla_cot(self, dataframe_a_mostrar):
        for item in self.tabla_cot_pos.get_children(): self.tabla_cot_pos.delete(item)
        if dataframe_a_mostrar is not None:
            df_limpio = dataframe_a_mostrar[["CODIGO", "DESCRIPCION", "PRECIO VENTA"]].fillna("")
            for i, (indice, fila) in enumerate(df_limpio.iterrows()):
                valores = fila.tolist()
                try:
                    if valores[2] != "": valores[2] = f"${float(valores[2]):,.0f}"
                except: pass
                self.tabla_cot_pos.insert("", "end", iid=str(indice), values=valores)
                if i % 50 == 0:
                    self.update_idletasks()

    def buscar_producto_cot(self, event=None):
        if self.catalogo_memoria is None: return
        palabra = self.entrada_busqueda_cot.get().lower().strip()
        if palabra == "":
            self.actualizar_tabla_cot(self.catalogo_memoria)
            return
        palabra_limpia = ''.join(c for c in unicodedata.normalize('NFKD', palabra) if unicodedata.category(c) != 'Mn')
        filtro = self.catalogo_memoria['columna_busqueda'].astype(str).str.contains(palabra_limpia, na=False)
        self.actualizar_tabla_cot(self.catalogo_memoria[filtro])

    def agregar_al_carrito_cot(self, event=None):
        seleccion = self.tabla_cot_pos.selection()
        if not seleccion: return
        indice_prod = seleccion[0] 
        if indice_prod in self.carrito_cotizacion:
            self.carrito_cotizacion[indice_prod]["cantidad"] += 1
        else:
            datos = self.catalogo_memoria.loc[int(indice_prod)]
            self.carrito_cotizacion[indice_prod] = {
                "nombre": str(datos["DESCRIPCION"]),
                "precio": float(datos["PRECIO VENTA"]) if pd.notna(datos["PRECIO VENTA"]) else 0.0,
                "cantidad": 1
            }
        self.actualizar_vista_carrito_cot(indice_seleccionado=indice_prod)

    def modificar_cantidad_cot(self, accion):
        seleccion = self.tabla_carrito_cot.selection()
        if not seleccion: return
        indice = seleccion[0]
        if accion == "sumar": self.carrito_cotizacion[indice]["cantidad"] += 1
        elif accion == "restar":
            self.carrito_cotizacion[indice]["cantidad"] -= 1
            if self.carrito_cotizacion[indice]["cantidad"] <= 0:
                del self.carrito_cotizacion[indice]
                indice = None 
        elif accion == "quitar":
            del self.carrito_cotizacion[indice]
            indice = None 
        self.actualizar_vista_carrito_cot(indice_seleccionado=indice)
        
    def agregar_extra_cotizacion(self, desc, precio, cant):
        id_unico = f"EXTRA-{int(time.time())}"
        self.carrito_cotizacion[id_unico] = {
            "nombre": desc,
            "precio": float(precio),
            "cantidad": int(cant)
        }
        self.actualizar_vista_carrito_cot(indice_seleccionado=id_unico)

    def actualizar_vista_carrito_cot(self, indice_seleccionado=None):
        for item in self.tabla_carrito_cot.get_children(): self.tabla_carrito_cot.delete(item)
        total = 0.0
        for indice, item in self.carrito_cotizacion.items():
            subtotal = item["precio"] * item["cantidad"]
            total += subtotal
            self.tabla_carrito_cot.insert("", "end", iid=indice, values=(item["nombre"], item["cantidad"], f"${subtotal:,.0f}"))
            
        neto = total / 1.19
        iva = total - neto
        
        self.lbl_subtotal_cot.configure(text=f"Neto: $ {neto:,.0f}")
        self.lbl_iva_cot.configure(text=f"IVA (19%): $ {iva:,.0f}")
        self.lbl_total_cot.configure(text=f"TOTAL: $ {total:,.0f}")
        if indice_seleccionado and indice_seleccionado in self.carrito_cotizacion:
            self.tabla_carrito_cot.selection_set(indice_seleccionado)

    def limpiar_cotizacion(self):
        self.carrito_cotizacion.clear()
        self.actualizar_vista_carrito_cot()
        for ent in self.entradas_cliente.values():
            ent.delete(0, 'end')

    def abrir_ventana_extra_cot(self):
        VentanaProductoExtra(self, destino="cotizacion")

    # ==========================================
    # NUEVAS FUNCIONES PARA ÍTEM EXTRA EN VENTAS
    # ==========================================
    def abrir_ventana_extra_pos(self):
        VentanaProductoExtra(self, destino="pos")

    def agregar_extra_pos(self, desc, precio, cant):
     
        id_unico = f"EXTRA-{int(time.time())}"
        
        self.carrito[id_unico] = {
            "nombre": desc,
            "precio": float(precio),
            "cantidad": int(cant)
        }
        self.actualizar_vista_carrito(indice_seleccionado=id_unico)

    def generar_pdf_cotizacion(self):
        if not self.carrito_cotizacion:
            messagebox.showwarning("Aviso", "No hay productos en la cotización.")
            return
            
        cliente = {k: v.get().strip() for k, v in self.entradas_cliente.items()}
        if not cliente["Nombre / Institución *"]:
            messagebox.showwarning("Aviso", "El Nombre/Institución es obligatorio.")
            return

        corr = 1
        if os.path.exists("correlativo_cot.txt"):
            with open("correlativo_cot.txt", "r") as f:
                corr = int(f.read().strip())
        id_cot = f"COT-{corr:04d}"

        pdf = FPDF()
        pdf.add_page()
        
        if os.path.exists("logo.png"): pdf.image("logo.png", x=10, y=8, w=45)
        elif os.path.exists("logo.jpg"): pdf.image("logo.jpg", x=10, y=8, w=45)

        pdf.set_font("helvetica", "B", 18)
        pdf.cell(0, 10, "ELECTROMUSIC", new_x="LMARGIN", new_y="NEXT", align="R")
        
        pdf.set_font("helvetica", "", 10)
        pdf.cell(0, 5, "Mariela Baeza B. ", new_x="LMARGIN", new_y="NEXT", align="R")
        pdf.cell(0, 5, "Rut: 12.064.910-8 ", new_x="LMARGIN", new_y="NEXT", align="R")
        pdf.cell(0, 5, "Teléfono: +56937742995", new_x="LMARGIN", new_y="NEXT", align="R")
        pdf.cell(0, 5, "Merced 254-B, San Felipe, V región", new_x="LMARGIN", new_y="NEXT", align="R")
        pdf.cell(0, 5, "Asesor de venta: Alex Charrier", new_x="LMARGIN", new_y="NEXT", align="R")
        pdf.cell(0, 5, "Teléfono: +56928594538", new_x="LMARGIN", new_y="NEXT", align="R")
        pdf.cell(0, 5, "Email: electromusicsanfelipe@gmail.com", new_x="LMARGIN", new_y="NEXT", align="R")

        pdf.ln(3) 
        
        pdf.set_font("helvetica", "B", 12)
        pdf.cell(0, 6, "Cotización Comercial", new_x="LMARGIN", new_y="NEXT", align="R")
        pdf.set_font("helvetica", "", 11)
        pdf.cell(0, 6, f"Nº: {id_cot}", new_x="LMARGIN", new_y="NEXT", align="R")
        pdf.cell(0, 6, f"Fecha: {datetime.datetime.now().strftime('%d/%m/%Y')}", new_x="LMARGIN", new_y="NEXT", align="R")
        pdf.ln(10)

        pdf.set_font("helvetica", "B", 12)
        pdf.cell(0, 8, "Preparado para:", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("helvetica", "", 10)
        for k, v in cliente.items():
            k_limpio = k.replace(" *", "")
            if v: pdf.cell(0, 6, f"{k_limpio}: {v}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(10)

        pdf.set_font("helvetica", "B", 10)
        pdf.cell(15, 8, "Cant", border=1, align="C")
        pdf.cell(115, 8, "Descripcion", border=1)
        pdf.cell(30, 8, "P. Unitario", border=1, align="C")
        pdf.cell(30, 8, "Total", border=1, align="C", new_x="LMARGIN", new_y="NEXT")

        pdf.set_font("helvetica", "", 10)
        total_cot = 0
        for item in self.carrito_cotizacion.values():
            cant = item["cantidad"]
            desc = item["nombre"][:60] 
            precio = item["precio"]
            sub = cant * precio
            total_cot += sub
            
            pdf.cell(15, 8, str(cant), border=1, align="C")
            pdf.cell(115, 8, desc, border=1)
            pdf.cell(30, 8, f"${precio:,.0f}", border=1, align="R")
            pdf.cell(30, 8, f"${sub:,.0f}", border=1, align="R", new_x="LMARGIN", new_y="NEXT")

        neto = total_cot / 1.19
        iva = total_cot - neto
        pdf.ln(5)
        pdf.set_font("helvetica", "B", 10)
        pdf.cell(130, 8, "", border=0)
        pdf.cell(30, 8, "Neto:", border=0, align="R")
        pdf.cell(30, 8, f"${neto:,.0f}", border=1, align="R", new_x="LMARGIN", new_y="NEXT")
        
        pdf.cell(130, 8, "", border=0)
        pdf.cell(30, 8, "IVA (19%):", border=0, align="R")
        pdf.cell(30, 8, f"${iva:,.0f}", border=1, align="R", new_x="LMARGIN", new_y="NEXT")

        pdf.set_font("helvetica", "B", 12)
        pdf.cell(130, 10, "", border=0)
        pdf.cell(30, 10, "TOTAL:", border=0, align="R")
        pdf.cell(30, 10, f"${total_cot:,.0f}", border=1, align="R", new_x="LMARGIN", new_y="NEXT")

        pdf.ln(15) 
        pdf.set_font("helvetica", "B", 10)
        pdf.cell(0, 6, "Términos y Condiciones:", new_x="LMARGIN", new_y="NEXT")
        
        pdf.set_font("helvetica", "", 9) 
        pdf.cell(0, 5, "1. Valores IVA incluido.", new_x="LMARGIN", new_y="NEXT")
        pdf.cell(0, 5, "2. Al aceptar esta cotización, avisar con anticipación para comprobar la existencia de stock.", new_x="LMARGIN", new_y="NEXT")
        pdf.cell(0, 5, "3. Datos Bancarios: Mariela Baeza \n Banco Santander \n Cta Corriente 000081732078", new_x="LMARGIN", new_y="NEXT")
        
        os.makedirs("Cotizaciones_PDF", exist_ok=True)
        
        nombre_cliente_limpio = re.sub(r'[\\/*?:"<>|]', "", cliente['Nombre / Institución *'])
        nombre_cliente_limpio = nombre_cliente_limpio.replace(' ', '_')
        
        nombre_pdf = f"Cotizaciones_PDF/{id_cot}_{nombre_cliente_limpio}.pdf"
        pdf.output(nombre_pdf)

        with open("correlativo_cot.txt", "w") as f:
            f.write(str(corr + 1))

        messagebox.showinfo("Éxito", f"¡Cotización generada exitosamente!\n\nSe ha guardado en la carpeta 'Cotizaciones_PDF'\nArchivo: {nombre_pdf}")
        self.limpiar_cotizacion()

if __name__ == "__main__":
    app = AppVentas()
    app.mainloop()